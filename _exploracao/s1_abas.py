import openpyxl
P = r"C:\Users\cairo\.kimi-code\sessions\wd_portal-do-irt-e_b99cab1ec6f9\session_5445835d-af58-466d-83ee-f46f894296ef\attachments\f_af1ab2b9-e3bb-4498-92da-e2a11141f23c-Reequilibrio_Tributario_v7_IRT-E_Setor_Publico.xlsx"
wb = openpyxl.load_workbook(P, data_only=False)
wbv = openpyxl.load_workbook(P, data_only=True)
for name in wb.sheetnames:
    ws = wb[name]; wsv = wbv[name]
    print(f"=== {name!r} dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column} state={ws.sheet_state}")
