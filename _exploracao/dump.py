# -*- coding: utf-8 -*-
import sys, openpyxl
from openpyxl.worksheet.formula import ArrayFormula
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
P = r"C:\Users\cairo\.kimi-code\sessions\wd_portal-do-irt-e_b99cab1ec6f9\session_5445835d-af58-466d-83ee-f46f894296ef\attachments\f_af1ab2b9-e3bb-4498-92da-e2a11141f23c-Reequilibrio_Tributario_v7_IRT-E_Setor_Publico.xlsx"
wb = openpyxl.load_workbook(P, data_only=False)
wbv = openpyxl.load_workbook(P, data_only=True)
name = sys.argv[1]
r0 = int(sys.argv[2]) if len(sys.argv)>2 else 1
r1 = int(sys.argv[3]) if len(sys.argv)>3 else 9999
ws = wb[name]; wsv = wbv[name]
for row in ws.iter_rows(min_row=r0, max_row=r1):
    for c in row:
        v = c.value
        if v is None: continue
        cv = wsv[c.coordinate].value
        if isinstance(v, ArrayFormula):
            print(f"{c.coordinate}\tAF:{v.text}\tV:{cv!r}")
        elif isinstance(v, str) and v.startswith("="):
            print(f"{c.coordinate}\tF:{v}\tV:{cv!r}")
        else:
            print(f"{c.coordinate}\t{v!r}")
